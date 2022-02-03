# -*- coding: utf-8 -*-

import re
import ply.lex as lex
import ply.yacc as yacc
import sys
import traceback
import networkx as nx
import openpyxl
import os
import json
import sys
import inspect
from enum import Enum
from networkml.error import NetworkError, NetworkMethodError, NetworkParseError, NetworkNotImplementationError
from networkml.error import NetworkReferenceError
from networkml.error import NetworkScriptInterruptionException
from networkml.generic import GenericCallable, GenericComponent, GenericValueHolder, Comparator, REPattern
from networkml.generic import GenericValidator, GenericDescription
from networkml.generic import debug, is_debug_mode, is_traceback
import networkml.genericutils as GU
from networkml.validator import GenericEvaluatee, BinaryEvaluatee, UnaryEvaluatee, GenericValidatorParam, TrueEvaluatee
import networkml.interpretermanager as IM
from networkml import config
import log4p

_armer = None


def get_armer():
    global _armer
    return _armer


def set_armer(armer):
    global _armer
    _armer = armer


class NetworkComponent(GenericComponent):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, **kwargs):
        super().__init__(owner)


class NetworkReturnValue(GenericValueHolder):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, value, success=True, reasons=None, cancel=False):
        self._value = value
        self._success = success
        self._reasons = reasons
        self._cancel = cancel

    @property
    def value(self):
        return self._value

    @property
    def fail(self):
        return not self._success

    @property
    def success(self):
        return self._success

    @property
    def reasons(self):
        return self._reasons

    @property
    def cancel(self):
        return self._cancel

    def __repr__(self):
        repr = "request success:{}, cancel proceed:{}, reason why:{},\n".format(self.success, self.cancel, self.reasons)
        if type(self.value) is list:
            repr = "{}value:[".format(repr)
            for e in self.value:
                repr = "{}\n{}".format(repr, e)
            repr = "{}]".format(repr)
            pass
        else:
            repr = "{}value:{}".format(repr, self.value)
        return repr


class DefaultSorter:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, sortee):
        super().__init__()
        self._sortee = sortee
        self._default_comparator = Comparator()

    @property
    def sortee(self):
        return self._sortee

    def minimum(self, sortee, comparator: Comparator = None):
        if comparator is None:
            comparator = self._default_comparator
        m = sortee[0]
        idx = 0
        for i in range(1, len(sortee)):
            if comparator.ordered(sortee[i], m):
                m = sortee[i]
                idx = i
        return m, idx

    def sort(self, comparator: Comparator = None):
        if comparator is None:
            comparator = self._default_comparator
        result = []
        target = self.sortee.copy()
        while len(target) != 0:
            m, idx = self.minimum(target, comparator)
            target.pop(idx)
            result.append(m)
        return result


class Literal:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, value):
        super().__init__()
        self._value = value

    @property
    def value(self):
        return self._value

    def __repr__(self):
        return "{}".format(self._value)


class Interval:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, start: int, end=None):
        if end is None:
            self._start = start
            self._end = end
        else:
            if start < end:
                self._start = start
                self._end = end
            else:
                self._start = end
                self._end = start

    @property
    def infinte(self) -> bool:
        return self._end is None

    @property
    def start(self) -> int:
        return self._start

    @property
    def end(self):
        return self._end

    def minimum(self):
        return self.start

    def maximum(self):
        return self.end

    def infinitize(self):
        self._end = None

    def contains(self, n):
        if n < self._start:
            return False
        elif self.infinte:
            return True
        elif self._end < n:
            return False
        return True

    def term(self) -> str:
        return self.__repr__()

    def __repr__(self):
        if self.infinte:
            m = ""
        else:
            m = self.end
        return "{}..{}".format(self.start, m)


class Numberset:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, *numbers):
        self._infinite = False
        self._intervals = []
        for n in numbers:
            self.add_number(n)

    @property
    def intervals(self) -> list:
        return self._intervals

    @property
    def infinite(self) -> bool:
        for i in self.intervals:
            if i.infinte:
                return True
        return False

    def infinitize(self):
        if self.infinite:
            return
        m = None
        for i in self._intervals:
            if m is None:
                m = i
            else:
                if m.maximum < i.maximum:
                    m = i
        if m is None:
            self.add_interval(Interval(1, None))
        else:
            m.infinitize()

    def add_number(self, n: int):
        for i in self._intervals:
            if i.contains(n):
                return
        self._intervals.append(Interval(n, n))

    def add_interval(self, i: Interval):
        removal = []
        start_hook = None
        end_hook = None
        for it in self._intervals:
            if it.start <= i.start and i.end <= it.end:
                return
            elif i.start <= it.start and it.end <= i.end:
                removal.append(it)
            elif it.start < i.start and it.end < i.end:
                start_hook = it
            elif i.start < it.start and i.end < it.end:
                end_hook = it
        for r in removal:
            self._intervals.remove(r)
        if start_hook is None and end_hook is None:
            self._intervals.append(i)
        elif start_hook is not None and end_hook is None:
            new_start = start_hook.start
            new_end = i.end
            new_interval = Interval(new_start, new_end)
            self._intervals.remove(start_hook)
            self._intervals.append(new_interval)
        elif start_hook is None and end_hook is not None:
            new_start = i.start
            new_end = end_hook.end
            new_interval = Interval(new_start, new_end)
            self._intervals.remove(start_hook)
            self._intervals.append(new_interval)

    @property
    def minimum(self):
        if self.infinite:
            return 1
        m = None
        for i in self._intervals:
            if m is None:
                m = i.start
            elif i.start < m:
                m = i.start
        return m

    @property
    def maximum(self):
        if self.infinite:
            return None
        m = None
        for i in self._intervals:
            if m is None:
                m = i.end
            elif m < i.end:
                m = i.end
        return m

    def contains(self, n: int):
        for i in self._intervals:
            if i.contains(n):
                return True
        return False

    def succ(self, n):
        if n == self.maximum:
            return None
        m = n+1
        while not self.contains(m):
            m += 1
        return m

    def term(self) -> str:
        return self.__repr__()

    def __repr__(self):
        series = None
        for i in self._intervals:
            if series is None:
                series = "{}".format(i)
            else:
                series = "{},{}".format(series, i)
        return "{" + series + "}"


class NumbersetOperator:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self):
        pass

    def multiply(self, N: Numberset, x) -> Numberset:
        new_N = None
        for i in N.intervals:
            new_i = Interval(i.start * x, i.end.x)
            if new_N is None:
                new_N = Numberset(new_i.start)
            new_N.add_interval(new_i)
        return new_N

    def product(self, N: Numberset, M: Numberset):
        new_N = None
        for n in N.intervals:
            for m in M.intervals:
                new_i = Interval(n.start * m.start, n.end * m.end)
                if new_N is None:
                    new_N = Numberset(new_i.start)
                new_N.add_interval(new_i)
        return new_N


class CommandOption:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, name: str, value=None, has_assignee=False, symbolic_assignee=False):
        self._name = name
        self._value = value
        self._has_assignee = has_assignee
        self._symbolic_assignee = symbolic_assignee

    @property
    def name(self):
        return self._name.replace("-", "")

    @property
    def value(self):
        return self._value

    @property
    def has_assignee(self):
        return self._has_assignee

    @property
    def symbolic_assignee(self):
        return self._symbolic_assignee

    def parse_value(self):
        pattern = r"[a-z]+=(?P<symbol>[a-z]+)|(?P<number>[0-9]+)|(?P<numberset>\{(\d+(,\d+)+)\})|(?P<anynumber>\{\d+\.\.\})"
        m = re.search(pattern, self.value)
        if m:
            dic = m.groupdict()
            if "symbol" in dic.keys():
                self._value = dic["symbol"]
            elif "number" in dic.keys():
                self._value = int(dic["number"])
            elif "numberset" in dic.keys():
                nums = dic["numberset"][1:len(dic["numberset"])-1].split(",")
                N = []
                for n in nums:
                    N.append(int(n))
                self._value = Numberset(N)
            elif "anynumber" in dic.keys():
                self._value = Numberset("*")
            else:
                raise NetworkError("unknown pattern:{}".format(self.value))
            pass
        else:
            raise NetworkError("unrecognized pattern:{}".format(self.value))

    def __repr__(self):
        if self.has_assignee:
            return "{}={}".format(self._name, self._value)
        else:
            return "{}".format(self._name)


class SimpleReach:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, forward: bool, edge_specs: GenericEvaluatee, node_specs: GenericEvaluatee):
        super().__init__()
        self._forward = forward
        self._edge_specs: GenericEvaluatee = edge_specs
        self._node_specs: GenericEvaluatee = node_specs

    @property
    def forward(self) -> bool:
        return self._forward

    @property
    def edge_specs(self) -> GenericEvaluatee:
        return self._edge_specs

    @property
    def node_specs(self) -> GenericEvaluatee:
        return self._node_specs

    def term(self) -> str:
        return self.__repr__()

    def __repr__(self):
        if self.forward:
            outbound = ">"
            inbound = ""
        else:
            outbound = ""
            inbound = "<"
        return "{}--[{}]--{}[{}]".format(inbound, self.edge_specs, outbound, self.node_specs)


class ReachNetwork:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, forward: bool, numbers: Numberset, reach: SimpleReach):
        self._forward = forward
        self._numbers = numbers
        self._reach = reach

    @property
    def numbers(self) -> Numberset:
        return self._numbers

    @property
    def reach(self) -> SimpleReach:
        return self._reach

    @property
    def forward(self) -> bool:
        return self._forward

    def term(self) -> str:
        return self.__repr__()

    def __repr__(self):
        return "({}){}".format(self.reach, self.numbers)


class ConstructiveReachNetwork:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, quantifier: Numberset, l, r=None, atomic: bool = False):
        self._quantifier = quantifier
        self._reaches = [l, r]
        self._atomic = atomic

    @property
    def quantifier(self) -> Numberset:
        return self._quantifier

    @property
    def reaches(self):
        return self._reaches

    @property
    def atomic(self) -> bool:
        return self._atomic

    @property
    def left(self):
        return self._reaches[0]

    @property
    def right(self):
        return self._reaches[1]

    def multiply(self, numbers: Numberset):
        operator = NumbersetOperator()
        self._quantifier = operator.product(self._quantifier, numbers)

    def serialize_left(self, most_left):
        if self.atomic:
            left: SimpleReach = self.left
            edge_connection = [most_left, left.edge_specs, left.node_specs]
            right_node_specs = left.node_specs
        else:
            edge_connection, right_node_specs = self.left.serialize_left(most_left)
        return edge_connection, right_node_specs

    def serialize_right(self, most_left):
        if self.right is None:
            return None, None
        if self.atomic:
            right: SimpleReach = self.right
            edge_connection = [most_left, right.edge_specs, right.node_specs]
            right_node_specs = right.node_specs
        else:
            edge_connection, right_node_specs = self.right.serialize_right(most_left)
        return edge_connection, right_node_specs

    def serialize(self, most_left):
        if self.atomic:
            left: SimpleReach = self.left
            edge_connection = [most_left, left.edge_specs, left.node_specs, self.quantifier]
            return edge_connection, left.node_specs
        else:
            left_edge_connection, left_right_node_specs = self.left.serialize(most_left)
            if self.right is None:
                return left_edge_connection
            else:
                right_edge_connection, right_right_node_spec = self.right.serialize(left_right_node_specs)
                return [left_edge_connection, right_edge_connection, self.quantifier], right_right_node_spec

    def term(self) -> str:
        return self.__repr__()

    def __repr__(self):
        if self.reaches[1] is None:
            return "({}){}".format(self.reaches[0], self.quantifier)
        else:
            return "(({}){}){}".format(self.reaches[0], self.reaches[1], self.quantifier)


class ReachNetworkConstructor:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self):
        pass

    def simple_reach(self, forward: bool, edge_specs: GenericEvaluatee, node_specs: GenericEvaluatee, numbers: Numberset):
        simple = SimpleReach(forward, edge_specs, node_specs)
        n = ConstructiveReachNetwork(numbers, simple, atomic=True)
        return n

    def extend(self, l: ConstructiveReachNetwork, r: ConstructiveReachNetwork, quantifier: Numberset):
        n = ConstructiveReachNetwork(quantifier, l, r)
        return n

    def multiply(self, n: ConstructiveReachNetwork, q: Numberset):
        n.multiply(q)
        return n


class ReachabilitySpecification(GenericComponent):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    # Constants
    SRC_SPEC = "src-spec"
    EDGE_SPEC = "edge-spec"
    DST_SPEC = "dst-spec"
    QUANTIFIER = "quantifier"
    DEPTHS = "depths"
    SRC_NODES = "src-nodes"
    LOOPBACK_NODES = "loopback-nodes"
    SRC_REACHABLE = "src-reachable"
    DST_NODES = "dst-nodes"
    DST_REACHABLE = "dst-reachable"

    def __init__(self, owner, node_specs: GenericEvaluatee):
        super().__init__(owner)
        self._node_specs = node_specs
        self._reaches = None
        self._sentence = None
        self._edge_dict = None
        self._serialization = None

    @property
    def serialization(self):
        return self._serialization

    @property
    def edge_dict(self):
        return self._edge_dict

    @edge_dict.setter
    def edge_dict(self, dic):
        self._edge_dict = dic

    @property
    def sentence(self) -> str:
        return self._sentence

    @sentence.setter
    def sentence(self, s):
        self._sentence = s

    def set_reaches(self, reaches: ConstructiveReachNetwork):
        self._reaches = reaches

    def serialize(self):
        if self.reaches is not None:
            edge_connection, _ = self.reaches.serialize(self.node_specs)
            self._serialization = edge_connection
            return edge_connection
            # print(edge_connection)

    def edge_dictionary(self, dic, c, binary_path, order):
        idx = binary_path.copy()
        idx.append(order)
        if len(c) == 4:  # simple edge
            src_spec = c[0]
            edge_spec = c[1]
            dst_spec = c[2]
            quantifier = c[3]
            index = "{}".format(idx[0])
            for i in idx[1:]:
                index = "{}.{}".format(index, i)
            dic[index] = {self.SRC_SPEC: src_spec,
                          self.EDGE_SPEC: edge_spec,
                          self.DST_SPEC: dst_spec,
                          self.QUANTIFIER: quantifier}
        else:  # compound:
            for i, d in enumerate(c[:len(c)-1]):
                self.edge_dictionary(dic, d, idx, i+1)
        self._edge_dict = dic

    @property
    def node_specs(self) -> GenericEvaluatee:
        return self._node_specs

    @property
    def reaches(self) -> ConstructiveReachNetwork:
        return self._reaches

    @property
    def outbound(self) -> bool:
        if self._reaches is None:
            return False
        elif self._reaches.forward:
            return True
        else:
            return False

    @property
    def inbound(self) -> bool:
        if self._reaches is None:
            return False
        elif self._reaches.forward:
            return False
        else:
            return True

    def term(self) -> str:
        return self.__repr__()

    def __repr__(self):
        return "[{}]{}".format(self.node_specs, self.reaches)

    @property
    def value(self):
        return True


class NetworkResolver(NetworkComponent):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner):
        super().__init__(owner)

    def get_method(self, caller, sig):
        raise NetworkNotImplementationError("get_method not implemented")


class NetworkArithmeticResolver(NetworkResolver):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    ATOMIC_RESOLVER = [(int, lambda x: x),
                       (bool, lambda x: x),
                       (str, lambda x: x),
                       (bool, lambda x: x)]
    ARITHMETIC_RESOLVERS = {
        "+": lambda c, a: a[0]+a[1],
        "-": lambda c, a: a[0]-a[1],
        "*": lambda c, a: a[0]*a[1],
        "%": lambda c, a: a[0]%a[1]
    }

    def __init__(self, owner):
        super().__init__(owner)

    def get_method(self, caller, sig):
        if sig in self.ARITHMETIC_RESOLVERS.keys():
            return self.ARITHMETIC_RESOLVERS[sig]
        return None

    def resolve_value(self, value):
        for r in self.ATOMIC_RESOLVER:
            if type(value) is r[0]:
                return r[1](value)


class NetworkResolverManager(NetworkResolver):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    RESOLVERS = []

    def __init__(self, owner):
        super().__init__(owner)

    def resolve_value(self, value):
        for r in self.RESOLVERS:
            if r.supported(value):
                return r.resolve_value(value)

    def append_resolver(self, resolver):
        self.RESOLVERS.append(resolver)


class NetworkSerializer:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    @property
    def opened(self):
        raise NetworkNotImplementationError("{}".format(self))

    def serialize(self, obj):
        raise NetworkNotImplementationError("{}".format(self))

    def deserialize(self):
        raise NetworkNotImplementationError("{}".format(self))

    def open(self, for_read=True, for_write=False):
        raise NetworkNotImplementationError("{}".format(self))

    def close(self):
        raise NetworkNotImplementationError("{}".format(self))


class NetworkTextSerializer(NetworkSerializer):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, filename):
        self._filename = filename
        self._mode = None
        self._opened = False
        self._f = None

    @property
    def filename(self):
        return self._filename

    @property
    def opened(self):
        return self._opened

    def serialize(self, obj):
        if self.opened and "w" in self._mode:
            self._f.write(obj)
        else:
            raise NetworkError("Invalid attempt for srialization")

    def deserialize(self):
        if self.opened and "r" in self._mode:
            return self._f.read()
        else:
            raise NetworkError("Invalid attempt for desrialization")

    def open(self, for_read=True, for_write=False):
        try:
            mode = ""
            if for_read:
                mode = "{}{}".format(mode, "r")
            if for_write:
                mode = "{}{}".format(mode, "w")
            if mode == "":
                raise NetworkError("Invalid attempt with options:{}".format(mode))
            self._opened = True
            self._f = open(self.filename, mode)
            self._mode = mode
        except Exception as ex:
            self._opened = False
            raise NetworkError("file {} open error.".format(self.filename))
        finally:
            pass

    def close(self):
        self._f = None
        self._opened = False


    #     self._signature = signature
    #     self._description = description
    #     self._script = script
    #     self._note = note
    #     pass
    #
    # @property
    # def signature(self):
    #     return self._signature
    #
    # @property
    # def description(self):
    #     return self._description
    #
    # def set_description(self, desc):
    #     self._description = desc
    #
    # @property
    # def script(self):
    #     return self._script
    #
    # def set_script(self, script):
    #     # print(script)
    #     self._script = script
    #
    # @property
    # def note(self):
    #     return self._note
    #
    # def set_note(self, note):
    #     self._note = note
    #
    # def analyze(self):
    #     pass
    #
    # def serialize(self, serializer):
    #     pass
    #
    # def deserialize(self, serializer):
    #     pass


# FIXME well-consider adequate interface for managing and representing document.
class NetworkDocument:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, obj, *args, **kwargs):
        self._obj = obj

    @property
    def desc(self):
        return None


# FIXME currently, no implementaion provided.
#       document management feature is one of important features in this framework.
class NetworkDocumentable:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    # FIXME doucment representation is not fixed.
    @property
    def document(self):
        return None


class NetworkBaseCallable(NetworkComponent, GenericCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    Lasted = 0
    ControlBroken = 1
    ControlReturned = 2
    REPORT_BREAK_POINT = "report_break_point"

    def __init__(self, owner, args=(), closer=False, cancel_stacking=True, safe_call=False, **kwargs):
        super().__init__(owner, args=args)
        self._args = args
        self._closer = closer
        self._cancel_stacking = cancel_stacking
        self._safe_call = safe_call

    @property
    def args(self):
        return self._args

    @property
    def closer(self):
        return self._closer

    @closer.setter
    def closer(self, _closer):
        self._closer = _closer

    @property
    def cancel_stacking(self):
        return self._cancel_stacking

    # def set_owners(self, caller, args):
    #     for a in args:
    #         if isinstance(a, NetworkVariable):
    #             a.set_owner(caller)

    def callee_type(self):
        return self

    @property
    def safe_call(self):
        return self._safe_call

    def evaluate(self, *args, **kwargs):
        return self(*args, **kwargs)

    def __call__(self, *args, **kwargs):
        if len(args) == 0:
            raise NetworkNotImplementationError("{}.__call__() without argument.".format(type(self)))
        caller = args[0]
        if not caller.running:
            raise NetworkScriptInterruptionException("script interrupted")
        if len(args) == 1:
            args = ()
        else:
            args = args[1]
        report_break_point = True
        if self.REPORT_BREAK_POINT in kwargs.keys():
            report_break_point = kwargs[self.REPORT_BREAK_POINT]
        stack_id = None
        st = caller.deepest_stack_id(caller)
        self.log.debug("**** PRE  STACK[{}] VARS:{}".format(st, caller.get_stack(caller, st)[caller.VARS]))
        if not self.cancel_stacking and not isinstance(caller, NetworkClassInstance):
            stack_id = caller.push_stack(caller)
        try:
            callee, actual_caller, actual_args = self.pre_call_impl(caller, args)
            # if caller != actual_caller:
            #     print("caller {} swapped to {}".format(caller, actual_caller))
            # if args != actual_args:
            #     print("args {} swapped to {}".format("(..{})".format(len(args)), "(..{})".format(len(actual_args))))
            # print("Method {}({}) running...".format(callee, (actual_caller, actual_args)))
            ret = callee(actual_caller, actual_args)
            self.log.debug("{}({})".format(callee, actual_args))
            # print("done.")
            return ret
        except NetworkScriptInterruptionException as ex:
            raise ex
        except Exception as ex:
            if is_traceback():
                self.log.debug("Tracebacking...")
                self.log.debug(traceback.format_exc())
            raise NetworkError("{}{}({}, {}) call failed.".format(type(self), self, caller, args), ex)
        finally:
            break_info = caller.break_point
            if not self.cancel_stacking and not isinstance(caller, NetworkClassInstance):
                caller.pop_stack(caller, stack_id)
            st = caller.deepest_stack_id(caller)
            self.log.debug("**** POST STACK[{}] VARS:{}".format(st, caller.get_stack(caller, st)[caller.VARS]))
            if report_break_point:
                caller.set_break_point(caller, break_info)

    def pre_call_impl(self, caller, args):
        # sub class must implement this method.
        raise NetworkNotImplementationError("{}.pre_call_impl not implemented.".format(self))


class NetworkCallable(NetworkBaseCallable, GenericCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, args=(), closer=False, cancel_stacking=True, safe_call=False, **kwargs):
        super().__init__(owner, args=args, close=closer, cancel_stacking=cancel_stacking, safe_call=safe_call)

    def do_nothing(self, caller, args):
        pass

    def pre_call_impl(self, caller, args):
        return self.call_impl, caller, self.args_impl(caller, args)

    def args_impl(self, caller, args, **kwargs):
        new_args = []
        for a in args:
            # print("** pre  ** ", type(a), a, "of a", type(caller), caller)
            # print("** post ** ", type(a), a, "of a", type(caller), caller)
            if isinstance(a, NetworkInstance):
                x = a
            elif isinstance(a, NetworkSymbol):
                x = caller.accessor.get(caller, a.symbol)
            elif isinstance(a, NetworkCallable):
                x = a(caller)
            elif isinstance(a, GenericEvaluatee):
                x = a.evaluate(caller)
            elif isinstance(a, GenericValueHolder):
                x = a.value
            else:
                x = a
            new_args.append(x)
        return tuple(new_args)

    def call_impl(self, caller, args, **kwargs):
        self.do_nothing(caller, args)


class NetworkInvoker:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def get_method(self, caller, m):
        raise NetworkNotImplementationError("Not implemented")

    def invoke(self, sig, caller, args):
        raise NetworkNotImplementationError("Not implemented")


# FIXME this class no longer needed.
# class HierachicalHolder(NetworkComponent):
#
#     def __init__(self, owner):
#         super().__init__(owner)
#         self._member_dict = {}
#         self._name_duplicated = False
#         self._indices = []
#
#     def set_owner(self, owner):
#         super().set_owner(owner)
#         for m in self._member_dict.values():
#             m.set_owner(owner)
#
#     @property
#     def name_duplicated(self):
#         return self._name_duplicated
#
#     def append_member(self, name, member):
#         if name in self._member_dict.keys():
#             debug("Name duplicated:{}".format(name))
#             self._name_duplicated = True
#         self._member_dict[name] = member
#
#     def append_index(self, index):
#         self._indices.append(index)
#
#     def has_member(self, name):
#         return name in self._member_dict.keys()
#
#     def get_member(self, name):
#         if name in self._member_dict.keys():
#             return self._member_dict[name]
#         else:
#             return None


# class NetworkMethodDocument(NetworkDocument):
#
#     def __init__(self, sig, method, *args, **kwargs):
#         super().__init__(sig)
#         self._method = method
#         self._arg_types = []
#         self._args = args
#         self._kwargs = kwargs
#         self._steps = []
#         self._callable = True
#         self._args_requirement = ""
#
#     @property
#     def args_requirement(self):
#         return self._args_requirement
#
#     def set_args_requirement(self, req):
#         self._args_requirement = req
#
#     @property
#     def callable(self):
#         return self._callable
#
#     @property
#     def arg_count(self):
#         return len(self._arg_types)
#
#     @property
#     def arg_types(self):
#         return self._arg_types
#
#     @property
#     def method(self):
#         return self._method
#
#     def analyze(self):
#         # FIXME implement method analysis statement(s) here
#         # argument analysis
#         for a in self.method.args:
#             self._arg_types.append(type(a))
#         # arg name duplication
#         args = []
#         for a in self.method.args:
#             if a not in args:
#                 args.append(a)
#         if len(args) == len(self.method.args):
#             self._callable = True
#         # process information
#         for i, s in enumerate(self.method.callees):
#             self._steps.append((i, s.callee_type))


class NetworkProperty(NetworkCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, name, writee=False, get_stmts=(), set_stmts=()):
        super().__init__(owner)
        self._name = name
        self._writee = writee
        self._set_stmts = set_stmts
        self._get_stmts = get_stmts

    def get(self, caller):
        rtn = None
        for s in self._get_stmts:
            rtn = s(caller)
        return rtn

    def set(self, caller, val):
        if self._writee:
            for s in self._set_stmts:
                rtn = s(caller)


class NetworkMethod(NetworkCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    class BreakPointState(Enum):

        Return = "return"
        Break = "break"
        Exception = "exception"

    def __init__(self, owner, signature, args=(), stmts=(), cancel_stacking=False, instance_method=False, *otherargs, **kwargs):
        super().__init__(owner, args=args, safe_call=True, cancel_stacking=cancel_stacking)
        self._signature = signature
        self._clazz = None
        self._instance_method = instance_method
        self._callees = []
        # self._document = NetworkMethodDocument(signature, self)
        # if "lexdata" in kwargs.keys():
        #     self._document.set_script(kwargs["lexdata"])
        self._globally = False
        if "globally" in kwargs.keys():
            self._globally = kwargs["globally"]
        for s in stmts:
            self._callees.append(s)

    @property
    def signature(self):
        return self._signature

    def set_owner(self, owner):
        super().set_owner(owner)
        for s in self._callees:
            s.set_owner(owner)

    @property
    def globally(self):
        return self._globally

    @property
    def is_global(self):
        return self.globally

    @property
    def clazz(self):
        if isinstance(self.owner, NetworkClassInstance):
            self._clazz = self.owner
        elif isinstance(self.owner, NetworkClass):
            self._clazz = self.owner
        else:
            self._clazz = None
        return self._clazz

    @clazz.setter
    def clazz(self, clazz):
        self._clazz = clazz

    @property
    def instance_method(self):
        return self._instance_method
    #@property
    #def document(self):
    #    return self._document

    @property
    def callable(self):
        return True
        # if self.document is None:
        #     return None
        # else:
        #     return self.document.callable

    @property
    def callees(self):
        return self._callees

    def append_callee(self, callee):
        return self._callees.append(callee)

    def actual_call_impl(self, caller, args, **kwargs):
        rtn = None
        for c in self.callees:
            rtn = c(caller)
            if isinstance(rtn, NetworkError):
                self.log.debug("Error Captured.")
                self.log.debug(rtn)
                self.log.debug("resumed")
                continue
            elif isinstance(rtn, NetworkReturnValue):
                self.log.debug("Special dealing with:")
                self.log.debug(rtn)
                continue
            if caller.break_point is not None:
                caller.set_break_point(caller, None)
                self.log.debug("caller.break_point: {0}".format(caller.break_point))
                if isinstance(rtn, NetworkReturnValue):
                    self.log.debug("Exception occurred, but safely resumed.")
                    self.log.debug(rtn)
                    return rtn.reasons
                elif isinstance(rtn, GenericValueHolder):
                    return rtn.value
                else:
                    return rtn
        return rtn

    def pre_call_impl(self, caller, args):
        # attach actual caller, if needed.
        # FIXME security controll should be implemented.
        return self.call_impl, caller, self.args_impl(caller, args)
        # No need to implement this method.

    def args_impl(self, caller, args, **kwargs):
        # FIXME argument adequateness varifiable here
        # new_args = []
        if len(args) < len(self.args):
            raise NetworkError("Method {} needs at least {} args, but {} given.".format(self, len(self.args), len(args)))
        args = super().args_impl(caller, args[:len(self.args)])
        for a, x in zip(self.args, args):
            caller.accessor.set(caller, a, x)
        return args

    def call_impl(self, caller, args, **kwargs):
        is_callable = self.callable
        if is_callable is not None and not is_callable:
            raise NetworkError("{}.{} is not in callable state.".format(self.clazz, self.signature))
        if is_callable is None:
            self.log.debug("Cannot analyze argumenet for {}.{}() ".format(self.owner, self.signature))
        if self.safe_call:
            try:
                return self.actual_call_impl(caller, args)
            except Exception as ex:
                self.log.debug("Internal Error occurred, but safely resumed.")
                self.log.debug(ex)
                return self.actual_call_impl(caller, args)
        else:
            return self.actual_call_impl(caller, args)

    def __repr__(self):
        args = "("
        if len(self.args) > 0:
            args = "{}{}".format(args, self.args[0])
            for a in self.args[1:]:
                args = "{}, {}".format(args, a)
        args = "{})".format(args)
        return "{}.{}{}".format(self.owner, self.signature, args)


# class NetworkInstanceDocument(NetworkDocument):
#
#     def __init__(self, sig, instance, clazz, *args, **kwargs):
#         super().__init__(sig)
#         self._instance = instance
#         self._clazz = clazz
#         self._args = args
#         self._kwargs = kwargs
#
#     @property
#     def instance(self):
#         return self._instance
#
#     @property
#     def clazz(self):
#         return self._clazz
#
#     def analyze(self):
#         # FIXME implement method analysis statement(s) here
#         # argument analysis
#         pass
#
#     def serialize(self, serializer):
#         serializer.serialize(self.script)
#         pass
#
#     def deserialize(self, serializer):
#         pass


class NetworkInstance(NetworkComponent, NetworkDocumentable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    PRIVATE = "private"
    PUBLIC = "public"
    PROTECTED = "protected"
    EMBEDDED = "embedded"
    STACK = "stack"
    VARS = "vars"
    METHODS = "methods"
    CLASSES = "classes"
    BREAK_POINT = "break_point"

    SELF = "$self"
    GENERATOR = "$generator"
    MANAGER = "$manager"
    OWNER = "$manager"

    def __init__(self, clazz, _id, owner, embedded=(), *args, **kwargs):
        super().__init__(owner)
        # Attribute area.
        self._attributes = {self.EMBEDDED: {},
                            self.PRIVATE: {},
                            self.PUBLIC: {},
                            self.PROTECTED: {},
                            self.STACK: [{self.VARS: {},
                                          self.METHODS: {},
                                          self.CLASSES: {},
                                          self.BREAK_POINT: None}]}
        #
        # Private attributes setting. Basically, object itself can access.
        self._attributes[self.PRIVATE][self.SELF] = self
        self._attributes[self.PRIVATE][self.GENERATOR] = clazz
        self._attributes[self.PRIVATE][self.MANAGER] = owner
        #
        # Adds properties for dynamic access. These are basically hidden from any object but self.
        for e in embedded:
            self._attributes[self.EMBEDDED][e[0]] = e[1]
            self.__dict__[e[0]] = e[1]
        #
        # print setting
        self._print_func = lambda x: self._print_buf.append(x)
        self._auto_flush = True
        self._print_buf = []
        #
        # FIXME deal with args
        #
        # FIXME validator and evaluator should be dynamically implemented.
        self._validator = None
        if "validator" in kwargs.keys():
            self._validator = kwargs["validator"]
        self._clazz = clazz
        self._id = _id
        #
        # FIXME accessor implementation should be dynamically changable.
        self._accessor = HierarchicalAccessor(self)
        self._enable_stack = False
        #
        # FIXME document managing feature is pended for further consideration.
        self._document = NetworkDocument(self)
        # self._document.set_script(lexdata)
        # FIXME this attribute should be removed
        self._globally = False
        if "globally" in kwargs.keys():
            self._globally = kwargs["globally"]
        self._running = False

    """
    Attribute accessors section.
    """
    # only for debug use
    def dump_attributes(self):
        att = "Attributes of {}:\n".format(self.signature)
        for i in range(len(self._attributes[self.STACK])):
            att += "stack[{}]\n".format(i)
            for k in self._attributes[self.STACK][i].keys():
                if k == "break_point":
                    continue
                att += "  {}\n".format(k)
                for f in self._attributes[self.STACK][i][k].keys():
                    att += "    {}\n".format(f)
        for k in self._attributes.keys():
            if k != self.STACK:
                att += k+"\n"
                for m in self._attributes[k].keys():
                    att += "  {}\n".format(m)
        return att

    # Caution! This method is internal method. Never call from other class scope.
    def _accessible_attributes(self, args=(), names=(), stack_criteria=None):
        acc = []
        for ac in args:
            if ac == self.STACK:
                stacks = self._attributes[self.STACK]
                i = len(stacks)-1
                while 0 <= i:
                    for c in stacks[i].keys():
                        if stack_criteria is not None and c not in stack_criteria:
                            continue
                        if c == self.BREAK_POINT:
                            continue
                        for k in stacks[i][c].keys():
                            if names is None or k in names:
                                acc.append((k, stacks[i][c][k]))
                    i = i-1
            elif ac in self._attributes.keys():
                for n in self._attributes[ac].keys():
                    if names is None or n in names:
                        acc.append((n, self._attributes[ac][n]))
            else:
                print("ERROR!!! ILLEGAL ACCESS ATTEMPTED!")
        return tuple(acc)

    # Caution! This method is internal method. Never call from other class scope.
    def _get_accessible_attributes(self, prior=(), names=()):
        att = self._accessible_attributes(prior, names)
        return att

    # Caution! This method is internal method. Never call from other class scope.
    def _has_accessible_attributes(self, prior=(), names=()):
        acc = self._accessible_attributes(prior, names)
        return len(acc) != 0

    # Caution! This method is internal method. Never call from other class scope.
    def _get_accessibilities(self, caller):
        # FIXME should consider caller.
        acc = []
        if caller == self:  # self is entirely accessible.
            acc.append(self.EMBEDDED)
            acc.append(self.PRIVATE)
            acc.append(self.PROTECTED)
            acc.append(self.PUBLIC)
            acc.append(self.STACK)
        elif caller == self.generator:  # generator is not accessible to PRIVATE and EMBEDDED.
            acc.append(self.PROTECTED)
            acc.append(self.PUBLIC)
            acc.append(self.STACK)
        elif caller == self.owner:  # owner is not accessible to PRIVATE, PROTECTED and EMBEDDED.
            acc.append(self.PUBLIC)
            acc.append(self.STACK)
        else:  # others are not accessible to PRIVATE and EMBEDDED. More protection should be implemented.
            acc.append(self.PUBLIC)
            acc.append(self.STACK)
        return tuple(acc)

    # This method is public accessible.
    def has_attribute(self, caller, name, security=None, globally=True):
        if security is None:
            security = (self.STACK, self.PUBLIC, self.PROTECTED, self.PRIVATE)
        acc = self._get_accessibilities(caller)
        _acc = []
        for a in acc:
            if a in security:
                _acc.append(a)
        att = self._accessible_attributes(_acc, names=[name])
        return len(att) != 0

    # This method is public accessible.
    def get_accessible_attribute(self, caller, name):
        acc = self._get_accessibilities(caller)
        att = self._get_accessible_attributes(acc, names=[name])
        return att[0][1]

    # This method is public accessible.
    def get_accessible_attributes(self, caller, name, identfier=None):
        acc = self._get_accessibilities(caller)
        att = self._get_accessible_attributes(acc, names=[name])
        rtn = []
        for a in att:
            b = a[1]
            if identfier is None:
                rtn.append(b)
            if identfier(b):
                rtn.append(b)
        return tuple(rtn)

    def get_attribute(self, caller, name):
        return self.get_accessible_attribute(caller, name)

    def get_attributes(self, caller, criteria=None):
        # FIXME accessibility selector 'criteria' should be implemented.
        acc = self._get_accessibilities(caller)
        att = self._accessible_attributes(acc, names=None)
        return att

    def _set_attribute(self, name, val, kind, depth=None, overwrite=None):
        """
        overwrite: None  -> set regardless var is set or not
                   True  -> set if var is set
                   False -> set if var is not set
        depth: if depth given, this method tries setting on depth or more deeper stack area
        """
        #print("_set_attribute", name, kind, depth, overwrite)
        # FIXME well-consider secure access.
        if isinstance(val, NetworkMethod):
            t = self.METHODS
        elif isinstance(val, NetworkClassInstance):
            t = self.CLASSES
        else:
            t = self.VARS
        if kind in (self.PRIVATE, self.PROTECTED, self.PUBLIC):
            if overwrite is None:
                self._attributes[kind][name] = val
                return
            if name not in self._attributes[kind].keys():
                if not overwrite:
                    self._attributes[kind][name] = val
                    return
            else:
                if overwrite:
                    self._attributes[kind][name] = val
                    return
                else:
                    # FIXME prohibition should be adequately notified.
                    raise NetworkError("Illegal attempt to set attribute. {}[{}] not set yet.".format(kind, name))
        elif kind in [self.STACK]:
            if depth is None:
                depth = self._deepest_stack_id()
            if overwrite is None:
                self._attributes[kind][depth][t][name] = val
                return
            while 0 <= depth:
                if name not in self._attributes[kind][depth][t].keys():
                    if not overwrite:
                        self._attributes[kind][depth][t][name] = val
                        return
                else:
                    if overwrite:
                        self._attributes[kind][depth][t][name] = val
                        return
                    else:
                        # FIXME prohibition should be adequately notified.
                        raise NetworkError("Illegal attempt to set attribute. {}[{}][{}] not set yet.".format(kind, t, name))
                depth = depth - 1
        # FIXME prohibition should be adequately notified.
        raise NetworkError("Illegal attempt to set attribute. {}[{}][{}] not set yet.".format(kind, t, name))

    def _remove_attribute(self, name, val, kind, depth=None):
        # FIXME well-consider secure access.
        if isinstance(val, NetworkCallable):
            t = self.METHODS
        elif isinstance(val, NetworkClassInstance):
            t = self.CLASSES
        else:
            t = self.VARS
        if kind in (self.PRIVATE, self.PROTECTED, self.PUBLIC):
            if name in self._attributes[kind].keys():
                self._attributes[kind].pop(name)
        elif kind in [self.STACK]:
            if depth is None:
                depth = self._deepest_stack_id()
            while 0 <= depth:
                if name in self._attributes[kind][depth][t].keys():
                    self._attributes[kind][depth][t].pop(name)
                depth = depth - 1
        # FIXME prohibition should be adequately notified.
        raise NetworkError("Illegal attempt to remove attribute. Not set yet.")

    def set_attribute(self, caller, name, val, kind=None, globally=False, depth=None, overwrite=None):
        # FIXME well-consider secure access.
        # print("***** WARNING!!! SET_ATTRIBUTE IS DEPRECATED.")
        if kind is None:
            kind = self.STACK
        if globally:
            depth = 0
        self._set_attribute(name, val, kind=kind, depth=depth, overwrite=overwrite)

    def remove_attribute(self, caller, name, val, kind=None, depth=None, overwrite=None):
        # FIXME well-consider secure access.
        # print("***** WARNING!!! REMOVE_ATTRIBUTE IS DEPRECATED.")
        if kind is None:
            kind = self.STACK
        self._set_attribute(name, val, kind=kind, depth=depth, overwrite=overwrite)

    def has_private_attribute(self, caller, name):
        # FIXME well-consider implementation of secure access.
        self.log.debug("***** WARNING!!! HAS_PRIVATE_ATTRIBUTE IS DEPRECATED.")
        att = self._get_accessible_attributes([self.PUBLIC], [name])
        return len(att) != 0

    def get_private_attribute(self, caller, name):
        # FIXME well-consider implementation of secure access.
        # print("***** WARNING!!! GET_PRIVATE_ATTRIBUTE IS DEPRECATED.")
        att = self._get_accessible_attributes([self.PUBLIC], [name])
        return att[0][1]

    def set_private_attribute(self, caller, name, val):
        # FIXME well-consider implementation of secure access.
        # print("***** WARNING!!! SET_PRIVATE_ATTRIBUTE IS DEPRECATED.")
        self._set_attribute(name, val, self.PUBLIC)

    # local variable
    def declare_var(self, var, globally=False):
        # FIXME this method is deprecated. Use push_stack instead.
        self.log.debug("***** WARNING!!! DECLARE_VAR IS DEPRECATED.")
        if globally:
            depth = 0
        else:
            depth = None
        self.set_attribute(self, var.name, var, kind=self.STACK, depth=depth, overwrite=None)
        var.set_owner(var)

    def register_var(self, caller, name, var, kind=None, depth=None):
        # FIXME well-consider secure access.
        if kind is None:
            kind = self.STACK
        self.set_attribute(caller, name, var, kind=kind, depth=depth)
        var.set_owner(self)

    def remove_var(self, caller, var, kind=None, depth=None):
        # FIXME well consider access security.
        if kind is None:
            kind = self.STACK
        self.remove_attribute(caller, var.name, var, kind=kind, depth=depth)

    # local method
    def declare_method(self, method, globally=False):
        # FIXME this method is deprecated. Use register_method instead.
        self.log.debug("***** WARNING!!! DECLARE_METHOD IS DEPRECATED.")
        if globally:
            depth = 0
        else:
            depth = None
        self.register_method(self, method.signature, method, kind=self.STACK, depth=depth, overwrite=None)

    # local method
    def register_method(self, caller, name, method, kind=None, depth=None, overwrite=None):
        # FIXME well-consider secure access.
        if kind is None:
            kind = self.STACK
        if depth is None:
            depth = 0
        self.set_attribute(caller, name, method, kind=kind, depth=depth, overwrite=overwrite)
        method.set_owner(self)

    def remove_method(self, caller, method, kind=None, depth=None, overwrite=None):
        # FIXME well consider access security.
        if kind is None:
            kind = self.STACK
        self.remove_attribute(caller, method.signature, method, kind=kind, depth=depth, overwrite=overwrite)

    def declare_class(self, clazz, globally=False):
        if clazz.globally or globally:
            depth = 0
        else:
            depth = None
        self.register_class(self, clazz.signature, clazz, kind=self.STACK, depth=depth, overwrite=None)
        clazz.set_parent(self)

    def register_class(self, caller, name, clazz, kind=None, depth=None, overwrite=None):
        # FIXME well-consider secure access.
        if kind is None:
            kind = self.STACK
        if depth is None:
            depth = 0
        self.set_attribute(caller, name, clazz, kind=kind, depth=depth, overwrite=overwrite)
        clazz.set_owner(self)

    def remove_class(self, caller, clazz, kind=None, depth=None):
        # FIXME well consider access security.
        if kind is None:
            kind = self.STACK
        self.remove_attribute(caller, clazz.signature, clazz, kind=kind, depth=depth)

    def get_callable(self, caller, sig):
        # FIXME well-consider access security, since this inspects private attributes.
        #if sig == "clazz":
        # self.log.debug("*** CALLER {0} {1}".format(type(caller), caller))
        # self.log.debug("*** Searching GET_METHOD")
        callee = self.get_method(caller, sig)
        if callee is not None:
            return callee
        # find callable attribute
        #if sig == "clazz":
        # self.log.debug("*** CALLER {0} {1}".format(type(caller), caller))
        # self.log.debug("*** Searching CALLABLE_ATTRIBUTE")
        weakest_acc = self._get_accessibilities(caller)
        att = self._accessible_attributes(weakest_acc, names=[sig], stack_criteria=(self.VARS, self.CLASSES))
        for a in att:
            if isinstance(a[1], NetworkCallable):
                return a[1]
        # if self.STACK in weakest_acc:
        #     att = self._accessible_attributes([self.STACK], names=[sig], stack_criteria=(self.VARS, self.CLASSES))
        #     for a in att:
        #         if isinstance(a[1], NetworkCallable):
        #             return a[1]
        return None

    def get_method(self, caller, sig):
        # This method is opened method, so basically arrows access to registered method, not callable variable.
        # search in secured area.
        weakest_acc = self._get_accessibilities(caller)
        att = self._accessible_attributes(weakest_acc, names=[sig], stack_criteria=[self.METHODS])
        if len(att) != 0:
            return att[0][1]
        # if self.STACK in weakest_acc:
        #     att = self._accessible_attributes([self.STACK], names=[sig], stack_criteria=[self.METHODS])
        #     if len(att) != 0:
        #         return att[0][1]
        m = self.clazz.get_method(caller, sig)
        if m is not None:
            return m
        if self.clazz.super_class is not None:
            m = self.clazz.super_class.get_method(caller, sig)
            if m is not None:
                return m
        # search generator method. This corresponds to call class method.
        # if self.generator is not None:
        #     att = self.generator.get_method(caller, sig)
        #     if att is not None:
        #         if isinstance(att, NetworkMethod):
        #             return att
        return None

    def get_class(self, caller, sig):
        # print("This method is updated for secure access.")
        # FIXME consider class hierarchy and security
        weakest_acc = self._get_accessibilities(caller)
        if self.STACK in weakest_acc:
            att = self._accessible_attributes([self.STACK], names=[sig], stack_criteria=[self.CLASSES])
            if len(att) != 0:
                return att[0][1]
        # att = self.get_accessible_attribute(caller, sig)
        # if att is not None:
        #     if isinstance(att, NetworkClassInstance):
        #         return att
        if self.generator is not None:
            att = self.generator.get_class(caller, sig)
            if att is not None:
                if isinstance(att, NetworkClassInstance):
                    return att
        # # search generator method.
        # if self.generator is not None:
        #     att = self.generator.get_class(caller, sig)
        #     if att is not None:
        #         if isinstance(att, NetworkClassInstance):
        #             return att
        return None

    def get_var(self, name, globally=True):
        # FIXME consider class hierarchy and security. currently, depends on self accessor method.
        # FIXME this method is deprecated.
        self.log.debug("***** WARNING!!! GET_VAR IS DEPRECATED.")
        self.log.debug("This method is deprecated. use accessor.get() instead.")
        if globally:
            depth = 0
        else:
            depth = None
        rtn = self.get_attribute(self, name)
        return rtn

    def set_var(self, var, globally=False):
        # FIXME this method is deprecated.
        self.log.debug("***** WARNING!!! SET_VAR IS DEPRECATED.")
        self.log.debug("This method is deprecated. use accessor.set() instead.")
        if globally:
            depth = 0
        else:
            depth = None
        self.set_attribute(self, var.name, var, self.STACK, depth=depth)

    @property
    def enable_stack(self):
        # FIXME secure control should be done.
        return self._enable_stack

    def set_stack_enable(self, en):
        self._enable_stack = en

    def _deepest_stack_id(self):
        return len(self._attributes[self.STACK])-1

    def deepest_stack_id(self, caller):
        # FIXME check accessibility
        return self._deepest_stack_id()

    def _get_stack(self, depth=None):
        if depth is None:
            depth = self._deepest_stack_id()
        # print("*** _GET_STACK with {} depth {}".format(self, depth))
        return self._attributes[self.STACK][depth]

    def get_stack(self, caller, depth=None):
        # FIXME check accessibility.
        return self._get_stack(depth)

    def _push_stack(self):
        frame = {self.VARS: {}, self.METHODS: {}, self.CLASSES: {}, self.BREAK_POINT: None}
        self._attributes[self.STACK].append(frame)
        return len(self._attributes[self.STACK])-1

    def push_stack(self, caller):
        # FIXME check accessibility.
        return self._push_stack()

    def _pop_stack(self, stack_id):
        if self.enable_stack:
            if stack_id is None:
                stack_id = self._deepest_stack_id()
            if self._deepest_stack_id() == 0:
                print("**** ILLEGAL _POP_STACK() ATTEMPTED!!!")
                return
            self._attributes[self.STACK] = self._attributes[self.STACK][0:stack_id]

    def pop_stack(self, caller, stack_id=None):
        # FIXME check accessibility.
        self._pop_stack(stack_id)

    @property
    def running(self) -> bool:
        return self._running

    def set_running(self, value: bool):
        self._running = value

    @property
    def context(self):
        # This method is deprecated. Use stack().
        return self.get_stack(self)

    def push_context(self):
        # FIXME this method is deprecated. Use push_stack instead.
        self.log.debug("***** WARNING!!! PUSH_CONTEXT IS DEPRECATED.")
        return self.push_stack(self)

    def pop_context(self, context_id):
        # FIXME this method is deprecated. Use push_stack instead.
        self.log.debug("***** WARNING!!! POP_CONTEXT IS DEPRECATED.")
        self.pop_stack(self, context_id)

    def get_accessible_attribute_map(self, caller, criteria=()):
        # FIXME secure access must be implemented right now.
        map = {}
        if caller == self:
            map[self.EMBEDDED] = self._attributes[self.EMBEDDED]
            map[self.PRIVATE] = self._attributes[self.PRIVATE]
        map[self.PROTECTED] = self._attributes[self.PROTECTED]
        map[self.PUBLIC] = self._attributes[self.PUBLIC]
        map[self.STACK] = self._attributes[self.STACK]
        return map

    def embedded_attrib(self, caller, args=()):
        # FIXME secure access should be well-considered.
        return tuple([_ for _ in self._attributes[self.EMBEDDED].keys()])

    def has_embedded_attrib(self, caller, attrib, args=()):
        # FIXME secure access should be well-considered.
        if caller == self:
            return attrib in self._attributes[self.EMBEDDED].keys()
        return False

    @property
    def clazz(self):
        return self._attributes[self.PRIVATE][self.GENERATOR]
        # if isinstance(self._clazz, NetworkClassInstance):
        #     return self._clazz
        # elif self.owner is not None:
        #     return self.owner.get_class(self.owner, self._clazz)
        # return None

    @property
    def generator(self):
        return self._attributes[self.PRIVATE][self.GENERATOR]
        #return self._clazz

    @property
    def id(self):
        return self._id

    @property
    def signature(self):
        if self.generator is None:
            return str(self.id)
        return "{}[{}]".format(self.generator, self.id)

    def set_owner(self, owner):
        super().set_owner(owner)
        self.set_attribute(self, self.MANAGER, owner, self.PUBLIC, overwrite=None)

    @property
    def parent(self):
        return self.owner

    def set_parent(self, parent):
        self.set_owner(parent)

    @property
    def globally(self):
        return self._globally

    @property
    def is_global(self):
        return self.globally

    @property
    def accessor(self):
        return self._accessor

    @property
    def document(self):
        return self._document

    @property
    def print_buf(self):
        return self._print_buf

    @property
    def auto_flush(self):
        return self._auto_flush

    @auto_flush.setter
    def auto_flush(self, auto):
        self._auto_flush = auto

    def set_print_func(self, func):
        self._print_func = func

    def print(self, *args):
        args = "".join([str(_) for _ in args])
        if self._auto_flush:
            self._print_func(args)
        else:
            self._print_buf.append(args)

    def flush_print_buf(self):
        buf = "\n".join([str(_) for _ in self._print_buf])
        self._print_buf.clear()
        return buf+"\n"

    def debug(self, obj, *args, **kwargs):
        obj = [str(obj)]
        obj.extend([str(_) for _ in args])
        _log = "log" in kwargs.keys() and not kwargs["log"]
        _stdout = "stdout" in kwargs.keys() and not kwargs["stdout"]
        _print_buf = "print_buf" in kwargs.keys() and not kwargs["print_buf"]
        _flush = "flush" in kwargs.keys() and not kwargs["flush"]
        if _log:
            self.log.debug(args=obj)
        if _stdout:
            for o in obj:
                sys.stdout.write(o)
            if _flush or self.auto_flush:
                sys.stdout.flush()
        else:
            if _flush or self.auto_flush:
                for o in obj:
                    self._print_func(o)
            else:
                for o in obj:
                    self._print_buf.append(o)

    @property
    def validator(self):
        if self._validator is not None:
            return self._validator
        return self.owner.validator

    def set_validator(self, v):
        self._validator = v

    def create_class(self, signature, owner, initializer_args=()):
        if signature in self.context["classes"].keys():
            raise NetworkError("Instance {} already has named class <{}>.".format(self, signature))
        #clazz = NetworkClassInstance(self, (signature, owner, initializer_args))  # local class doesn't has meta-class.
        clazz = NetworkClassInstance(self, signature, self, self.parent, embedded=(), *initializer_args)
        self.declare_class(clazz)
        return clazz

    def create_method(self, signature, args, stmt):
        if signature in self.context["methods"].keys():
            raise NetworkError("Instance {} already has named method '{}'.".format(self, signature))
        # m = NetworkMethod(self, (signature, args, stmt))
        m = NetworkMethod(self, signature, args=args, stmt=stmt)
        self.register_method(self, signature, m, depth=0, overwrite=None)
        return m

    @property
    def break_point(self):
        return self.get_stack(self)[self.BREAK_POINT]
        # return self.context["break_point"]

    def set_break_point(self, caller, break_info):
        stack = self.get_stack(caller)
        stack[self.BREAK_POINT] = break_info
        # self.context["break_point"] = break_info

    # def managing_methods(self):
    #     methods = []
    #     for i, ctx in enumerate(self._context):
    #         for k in ctx["methods"].keys():
    #             methods.append(ctx["methods"][k])
    #     return methods
    #
    # def managing_vars(self):
    #     vars = []
    #     for i, ctx in enumerate(self._context):
    #         for k in ctx["vars"].keys():
    #             vars.append(ctx["vars"][k])
    #     return vars
    #
    # def managing_classes(self):
    #     classes = []
    #     for i, ctx in enumerate(self._context):
    #         for k in ctx["classes"].keys():
    #             classes.append(ctx["classes"][k])
    #     return classes

    def invoke(self, caller, sig, args):
        # DO NOT ACCESS PRIVATE ATTRIBUTES.
        callee = self.get_method(caller, sig)
        if callee is None:
            # callee = self.get_callable(caller, sig)
            # if callee is None:
            raise NetworkError("Method {}.{} not found.".format(self.clazz.signature, sig))
        return callee(caller, args)

    def stack_structure(self):
        # FIXME DO NOT USE THIS METHOD!!! INSTEAD, USE _ATTRIBUTES().
        raise NetworkNotImplementationError("deprecated")
        # stack_structure = ""
        # for i, c in enumerate(self._context):
        #     vars = "vars: "
        #     for k in c["vars"].keys():
        #         vars = vars + k + ","
        #     methods = "methods: "
        #     for k in c["methods"].keys():
        #         methods = methods + k + ","
        #     classes = "classes: "
        #     for k in c["classes"].keys():
        #         classes = classes + k + ","
        #     stack_structure = stack_structure + "{}:[{}, {}, {}] ".format(i, vars, methods, classes)
        # return stack_structure

    def __repr__(self):
        return "{}".format(self.signature)


# class NetworkClassInstanceDocument(NetworkDocument):
#
#     def __init__(self, sig, instance, clazz, *args, **kwargs):
#         super().__init__(sig)
#         self._instance = instance
#         self._clazz = clazz
#         self._args = args
#         self._kwargs = kwargs
#
#     @property
#     def instance(self):
#         return self._instance
#
#     @property
#     def clazz(self):
#         return self._clazz
#
#     def analyze(self):
#         # FIXME implement method analysis statement(s) here
#         # argument analysis
#         pass
#
#     def serialize(self, serializer):
#         for i in self.clazz:
#             pass
#         serializer.serialize(self.script)
#         pass
#
#     def deserialize(self, serializer):
#         pass


class NetworkSymbol(UnaryEvaluatee):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, symbol):
        super().__init__(owner, symbol, symbol, as_symbol=True)
        self._symbol = symbol

    @property
    def symbol(self):
        return self._evaluatees[0]

    def set_symbol(self, sym):
        self._evaluatees[0] = sym

    def evaluate(self, caller=None):
        return super().evaluate(caller)

    def __repr__(self):
        return self.symbol


class NetworkClassInstance(NetworkInstance, NetworkCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, generator, signature, owner, super_class=None, embedded=(), *args, **kwargs):
        super().__init__(generator, signature, owner, embedded=embedded, args=args)
        self._super_class = super_class
        self._signature = signature
        self._globally = False
        #
        # FIXME class stored place is not defined in this timing.
        if "globally" in kwargs.keys():
            globally = kwargs["globally"]
        self._init_args = None
        if "init_args" in kwargs.keys():
            self._init_args = kwargs["init_args"]
        # self._lexdata = None
        # if "lexdata" in kwargs.keys():
        #     self._lexdata = kwargs["lexdata"]
        self._initializer = None
        self._last_instance = 0
        self._instance_ids = []
        # if super_class is not None:
        #     #sup = generator.get_class(self, super_class)
        #     sup = generator.accessor.get(generator, super_class, security=self.PUBLIC)
        #     print("*** GENERATOR:{}, SUPER_CLASS:{}, {}".format(generator, super_class, sup))
        #     self.accessor.set(self, super_class, sup, security=self.PUBLIC)
        #
        # FIXME document management feature is not implemented.
        # self._document = NetworkInstanceDocument(self.signature, self.id, self.clazz)
        # self._document.set_script(self._lexdata)
        # FIXME deal with init args

    @property
    def signature(self):
        # if self.super_class is None:
        #     return self._signature
        # return "{}::{}[{}]".format(self.super_class, self.generator, self._signature)
        return "{}".format(self._signature)

    @property
    def clazz(self):
        return self._attributes[self.PRIVATE][self.GENERATOR]
        #return self

    @property
    def super_class(self):
        if self._super_class is None:
            return None
        generator = self._attributes[self.PRIVATE][self.GENERATOR]
        sup = generator.accessor.get(generator, self._super_class, security=self.PUBLIC)
        return sup
        #return self.accessor.get(self, self._super_class, security=self.PUBLIC)

    @property
    def instance_ids(self):
        return self._instance_ids

    @property
    def next_instance_id(self):
        if len(self._instance_ids) == 0:
            return 1
        return self._instance_ids[len(self._instance_ids)-1] + 1

    def initialize_methods(self, initializer, method_list=()):
        initializer.set_owner(self)
        self._initializer = initializer
        initializer.set_owner(self)
        self.register_method(self, initializer.signature, initializer, depth=0, overwrite=None)
        self.log.debug(str(initializer))
        for e in method_list:
            e.set_owner(self)
            self.register_method(self, e.signature, e, depth=0, overwrite=None)
            self.log.debug(str(e))

    @property
    def globally(self):
        return self._globally

    @property
    def document(self):
        return self._document

    @property
    def is_global(self):
        return self.globally

    def push_stack(self, caller):
        self.log.debug("**** NetworkClassInstance doesn't support context stack operation.")
        return self.deepest_stack_id(caller)

    def pop_stack(self, caller, stack_id=None):
        self.log.debug("**** NetworkClassInstance doesn't support context stack operation.")
        pass

    def declare_var(self, var, globally=False):
        super().declare_var(var, globally=globally)

    def declare_method(self, method, globally=True):
        super().declare_method(method, globally=globally)

    def get_method(self, caller, sig):
        # FIXME consider class hierarchy and security
        weakest_acc = self._get_accessibilities(caller)
        att = self._accessible_attributes(weakest_acc, names=[sig], stack_criteria=[self.METHODS])
        if len(att) != 0:
            return att[0][1]
        # m = super().get_method(caller, sig)
        # if m is not None:
        #     return m
        if self.super_class is not None:
            m = self.super_class.get_method(caller, sig)
            if m is not None:
                return m
        return None

    def get_initializer(self) -> NetworkMethod:
        initializer = self._initializer
        return initializer

    def create_instance(self, _id, owner, args):
        # inherits dynamic attributes.
        embedded = []
        for e in self._attributes[self.EMBEDDED].keys():
            embedded.append((e, self._attributes[self.EMBEDDED][e]))
        embedded = tuple(embedded)
        # creates and initializes instance
        instance = NetworkInstance(self, _id, owner, embedded, args)
        initializer: NetworkMethod = self.get_initializer()
        instance.set_running(True)
        instance.set_stack_enable(False)
        if initializer is not None:
            actual_args = [instance]
            actual_args.extend(args)
            initializer(owner, actual_args)
        instance.set_stack_enable(True)
        return instance

    def release_instance(self, instance):
        if instance.clazz == self:
            self._instance_ids.remove(instance.id)
        else:
            debug("Not managed object {}".format(instance))

    def call_impl(self, caller, args, **kwargs):
        self._last_instance += 1
        self._instance_ids.append(self._last_instance)
        ret = self.create_instance(self._last_instance, caller, args)
        return ret

    def __repr__(self):
        if self._super_class is None:
            return self._signature
        return "{}::{}[{}]".format(self._super_class, self.generator, self._signature)


class NetworkClass(NetworkClassInstance):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, signature, super_class=None, embedded=(), globally=False):
        super().__init__(owner, signature, owner, super_class=super_class, embedded=embedded)
        self._globally = globally
        self._classes = {}

    @property
    def signature(self):
        return "{}:<MetaClass>".format(super().signature)

    @property
    def globally(self):
        return self._globally

    @property
    def is_global(self):
        return self.globally

    @property
    def classes(self):
        return self._classes

    def create_instance(self, clazz_sig, caller, args):
        if clazz_sig in self._classes.keys():
            raise NetworkError("class {} is already created.".format(clazz_sig))
        # expects embedded=args[0] and args=args[1:]
        embedded = args[0]
        args = args[1:]
        clazz = NetworkClassInstance(self, clazz_sig, caller, super_class=None, embedded=embedded, args=args)
        self.classes[clazz.signature] = clazz
        return clazz

    def call_impl(self, caller, args, **kwargs):
        clazz_sig = args[0]
        args = args[1:]  # expects embedded=args[0] and args=args[1:]
        # print("clazz_sig:{0}, args:{1}".format(clazz_sig, args))
        clazz = self.create_instance(clazz_sig, caller, args)
        return clazz


class NetworkWritee(NetworkComponent):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner):
        super().__init__(owner)

    def write(self, value):
        raise NetworkNotImplementationError("not implemented")


class NetworkVariable(NetworkWritee, GenericValueHolder):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, name, globally=False):
        super().__init__(owner)
        self._name = name
        self._value = None
        self._globally = globally

    @property
    def name(self):
        return self._name

    @property
    def globally(self):
        return self._globally

    @property
    def is_global(self):
        return self.globally

    def write(self, value):
        if self.owner is None:
            raise NetworkError("{} cannot handle writing operation.:{}={}".format(self, self.name, value))
        self._value = value

    @property
    def value(self):
        if self.owner is None:
            raise NetworkError("{} cannot resolve name. owner doesn't exist.:{}".format(self, self.name))
        else:
            return self._value

    def __repr__(self):
        return "{}".format(self._name)


class SimpleVariable(NetworkVariable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, name, value, globally=False):
        super().__init__(owner, name, globally)
        self._value = value

    @property
    def value(self):
        if self.owner is None:
            debug("Bug!!! value reference without declaration.")
            return self._value
        else:
            return self._value

    def write(self, value):
        self._value = value

    def __repr__(self):
        return "{}".format(self.name)


class NetworkNothing(NetworkError):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, args):
        super().__init__(args)


class HierarchicalAccessor(NetworkComponent):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner=None):
        super().__init__(owner)

    # def separate_name(self, caller, name):
    #     return self._separate_name(caller, name)
    #

    def _separate_name(self, caller, name):
        if isinstance(name, NetworkSymbol):
            name = name.symbol
        names = name.split(".")
        last_segment = names[len(names)-1]
        #sympat = r"\s*(?P<symbol>(\$|)[a-zA-Z_]+([a-zAZ0-9_\$]*[a-zAZ0-9]+)*)\s*"
        #sympat = r"(?P<symbol>(\$|)[a-zA-Z_]+[a-zAZ0-9_\$]*(|[a-zA-Z_]+[a-zAZ0-9_]*|\[(\"[^\"]*\"|[0-9]+)\]))"
        sympat = r"(?P<symbol>(\$|)[a-zA-Z_]+[a-zAZ0-9_\$]*)"
        m = re.match(sympat, last_segment)
        if m is None:
            raise NetworkReferenceError("Invalid reference format:{}".format(last_segment))
        last_name = m.groupdict()['symbol']
        indices_segment = last_segment[m.span()[1]:]
        indices = []
        idxpat = r"^\s*\[\s*((?P<number>\d+)|(?P<literal>(\"[^\"]+\"|'[^']+'))|(?P<symbol>[a-zA-Z_]+([a-zAZ0-9_]*[a-zAZ0-9]+)*))\s*\]"
        while True:
            m = re.match(idxpat, indices_segment)
            if m is None:
                break
            num_idx = m.groupdict()['number']
            sym_idx = m.groupdict()['symbol']
            ltr_idx = m.groupdict()['literal']
            if num_idx is not None and num_idx != "":
                indices.append(int(num_idx))
            elif sym_idx is not None and sym_idx != "":
                self.log.debug("****** sym_idx: {} {}".format(type(sym_idx), sym_idx))
                sym_idx = caller.get_attribute(caller, sym_idx)
                # sym_idx = caller.accessor.get(caller, sym_idx)
                indices.append(sym_idx)
            elif ltr_idx is not None and ltr_idx != "":
                indices.append(ltr_idx)
            else:
                raise NetworkNothing(name)
            indices_segment = indices_segment[m.span()[1]:]
        if len(names) == 1:
            first_name = None  # last_name
            middle_names = ()
            # last_name = None
        else:
            first_name = names[0]
            middle_names = names[1:len(names)-1]
        return first_name, middle_names, last_name, indices

    def get_indexed_value(self, caller, var, indices):
        for i in indices:
            if isinstance(i, int):
                if i < 0:
                    raise NetworkReferenceError("Negative indexer {} assigned.".format(i))
                if isinstance(var, GenericValueHolder):
                    var = var.value
                if not (isinstance(var, list) or isinstance(var, tuple)):
                    raise NetworkReferenceError("Unaccessible referee {} for index {}.".format(var, i))
                var = var[i]
            elif isinstance(i, str):
                if type(var) is dict:
                    var = var[i]
                else:
                    raise NetworkReferenceError("Unaccessible referee {} for index {}.".format(var, i))
        return var

    def _set_to_indexed_object(self, var, indices, val):
        i = indices[len(indices)-1]
        for i in indices[0:len(indices)-1]:
            if isinstance(i, int):
                if i < 0:
                    raise NetworkReferenceError("Negative indexer {} assigned.".format(i))
                if not isinstance(var, list):
                    raise NetworkReferenceError("Unaccessible referee {} for index {}.".format(var, i))
                else:
                    var = var[i]
            elif isinstance(i, str):
                if not isinstance(var, dict):
                    raise NetworkReferenceError("Unaccessible referee {} for index {}.".format(var, i))
                var = var[i]
            else:
                NetworkReferenceError("Unavailable indexer {}.".format(i))
        if isinstance(i, int):
            if i < 0:
                raise NetworkReferenceError("Negative indexer {} assigned.".format(i))
            if not isinstance(var, list):
                raise NetworkReferenceError("Unaccessible referee {} for index {}.".format(var, i))
            else:
                var[i] = val
        elif isinstance(i, str):
            if not isinstance(var, dict):
                raise NetworkReferenceError("Unaccessible referee {} for index {}.".format(var, i))
            var[i] = val
        else:
            return NetworkNothing("Unavailable indexer {}.".format(i))
        return True

    def _get_named_value(self, caller, first_name, middle_names, last_name, indices, security=NetworkInstance.STACK, globally=False):
        return self._get_named_object(caller, first_name, middle_names, last_name, indices, security=security, globally=globally)

    def _get_named_object(self, caller, first_name, middle_names, last_name, indices, security=NetworkInstance.STACK, globally=False):
        names = []
        if first_name is not None:
            names.append(first_name)
        names.extend(middle_names)
        if last_name is not None:
            names.append(last_name)
        obj = caller
        for i, n in enumerate(names):
            if isinstance(obj, NetworkInstance):
                if not obj.has_attribute(caller, n):
                    symbol = self.complex_name(names, ())
                    raise NetworkNothing("Couldn't access to {}.{}. '{}' not found.".format(obj, symbol, n))
                obj = obj.get_attribute(obj, n)
            elif isinstance(obj, dict):
                if n not in obj.keys():
                    symbol = self.complex_name(names, ())
                    raise NetworkNothing("Couldn't access to {}.{} to get due to unavailabel dict index.".format(obj, symbol))
                obj = obj[n]
            elif type(obj) is list or type(obj) is tuple:
                if type(n) is not int:
                    symbol = self.complex_name(names, ())
                    raise NetworkNothing("Couldn't access to {}.{} to get due to index isn't int.".format(obj, symbol))
                obj = obj[n]
            else:
                symbol = self.complex_name(names, ())
                raise NetworkNothing("Couldn't access to {}.{} to get anyway.".format(obj, symbol))
        for j, i in enumerate(indices):
            if isinstance(i, str):
                if (i[0] == "\"" and i[len(i) - 1] == "\"") or (i[0] == "'" and i[len(i) - 1] == "'"):
                    i = i[1:len(i) - 1]
                    if not isinstance(obj, dict):
                        symbol = self.complex_name(names, indices[:j + 1])
                        raise NetworkNothing("Couldn't access to {}.{} to get {} {} th index.".format(obj, symbol, type(i), j))
                    obj = obj[i]
                else:
                    i = self.get(caller, i, types=(int, str))
                    if type(i) is int and (type(obj) is list or type(obj) is tuple):
                        obj = obj[i]
                    elif type(i) is str and type(obj) is dict:
                        obj = obj[i]
                    else:
                        symbol = self.complex_name(names, indices[:j + 1])
                        raise NetworkNothing("Couldn't access to {}.{} to get {} {} th inddex.".format(obj, symbol, type(i), j))
            elif type(i) is int and (isinstance(obj, list) or isinstance(obj, tuple)):
                obj = obj[i]
            else:
                symbol = self.complex_name(names, indices[:j+1])
                raise NetworkNothing("Couldn't access to {}.{} to get {} {} th index.".format(obj, symbol, type(i), j))

        return obj

    def complex_name(self, names, indices):
        name = ".".join(names)
        idx = ""
        for j in indices:
            idx = "{}[{}]".format(idx, j)
        return "{}{}".format(name, idx)

    def get(self, caller, hierarchical_name, types=None, security=NetworkInstance.STACK, globally=False):
        original_caller = caller
        first_name, middle_names, last_name, indices = self._separate_name(caller, hierarchical_name)
        rtn = self._get_named_value(caller, first_name, middle_names, last_name, indices, security=security, globally=globally)
        if isinstance(rtn, NetworkNothing):
            raise rtn
        if not (type(types) is list or type(types) is tuple):
            available_types = [types]
        else:
            available_types = types
        for t in available_types:
            if t is None:
                return rtn
            if isinstance(rtn, t):
                return rtn
        raise NetworkNothing("{} not found.".format(hierarchical_name))

    def _set_named_value(self, caller, first_name, middle_names, last_name, indices, val, security=NetworkInstance.STACK,
                         globally=False, overwrite=None):
        return self._set_named_object(caller, first_name, middle_names, last_name, indices, val, security, globally)

    def _set_named_object(self, caller, first_name, middle_names, last_name, indices, val, security=NetworkInstance.STACK,
                          globally=False, overwrite=None):
        #print("_set_named_object", first_name, middle_names, last_name, indices, security, globally)
        names = []
        if first_name is not None:
            names.append(first_name)
        names.extend(middle_names)
        if last_name is not None:
            names.append(last_name)
        obj = caller
        if len(indices) == 0:
            last = len(names)-1
            last_something = names[last]
        else:
            last = len(names)
            last_something = indices[len(indices)-1]
        for i, n in enumerate(names[:last]):
            if isinstance(obj, NetworkInstance):
                if not obj.has_attribute(caller, n):
                    symbol = self.complex_name(names, ())
                    raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
                obj = obj.get_attribute(obj, n)
            elif isinstance(obj, dict):
                if n not in obj.keys():
                    symbol = self.complex_name(names, ())
                    raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
                obj = obj[n]
            elif i != len(names):
                symbol = self.complex_name(names, ())
                raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
        if len(indices) != 0:
            for j, ix in enumerate(indices[:len(indices)-1]):
                if isinstance(ix, str):
                    if ix[0] == "\"" and ix[len(ix) - 1] == "\"":
                        ix = ix[1:len(ix) - 1]
                    #     if not isinstance(obj, dict):
                    #         symbol = self.complex_name(names, indices[:j + 1])
                    #         raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
                    #     obj = obj[i]
                    else:
                        ix = self.get(caller, ix, types=(int, str))
                    # if type(i) is int and (type(obj) is list or type(obj) is tuple):
                    #     obj = obj[i]
                    # elif type(i) is str and type(obj) is dict:
                    if isinstance(ix, str) and isinstance(obj, dict):
                        obj = obj[ix]
                    elif isinstance(ix, int) and (isinstance(obj, list) or isinstance(obj, tuple)):
                        obj = obj[ix]
                    else:
                        symbol = self.complex_name(names, indices[:j + 1])
                        raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
                elif isinstance(ix, int) and (isinstance(obj, list) or isinstance(obj, tuple)):
                    obj = obj[ix]
                else:
                    symbol = self.complex_name(names, indices[:j + 1])
                    raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
        if isinstance(obj, NetworkInstance):
            #obj.set_attribute(obj, last_something, val)
            obj.set_attribute(caller, last_something, val, kind=security, globally=globally, overwrite=overwrite)
            return True
        i = last_something
        if isinstance(i, str):
            if last_something[0] == "\"" and i[len(i) - 1] == "\"": # literal
                i = i[1:len(i) - 1]
                # if not isinstance(obj, dict):
                #     symbol = self.complex_name(names, indices)
                #     raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
                # obj[i] = val
            else:
                i = self.get(caller, i, types=(int, str))
            # if type(i) is int and (type(obj) is list or type(obj) is tuple):
            #     obj[i] = val
            # elif type(i) is str and type(obj) is dict:
            if isinstance(i, str) and isinstance(obj, dict):
                obj[i] = val
            elif isinstance(i, int) and isinstance(obj, list):
                obj[i] = val
            else:
                symbol = self.complex_name(names, indices)
                raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
        elif isinstance(i, int):
            obj[i] = val
        else:
            symbol = self.complex_name(names, indices)
            raise NetworkNothing("Couldn't access to {}.{} to set.".format(obj, symbol))
        return True

    def set(self, caller, hierarchical_name, val, security=NetworkInstance.STACK, globally=False,
            overwrite=None, value=True, method=False, clazz=False):
        original_caller = caller
        first_name, middle_names, last_name, indices = self._separate_name(caller, hierarchical_name)
        ret = self._set_named_value(caller, first_name, middle_names, last_name, indices, val, security=security,
                                    globally=globally, overwrite=overwrite)
        return ret


class NetworkMethodCaller(NetworkCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, hierarchy_name, args=()):
        super().__init__(owner, args=args, cancel_stacking=True)
        name = hierarchy_name.symbol
        self._hierarchical_name = name
        names = name.split(".")
        self._actual_caller = ".".join(names[0:len(names)-1])
        self._method_name = names[len(names)-1]
        self._is_default_method = self._method_name == name

    # def set_owner(self, owner):
    #     super().set_owner(owner)

    @property
    def hierarchical_name(self):
        return self._hierarchical_name

    @property
    def actual_caller(self):
        return self._actual_caller

    @property
    def method_name(self):
        return self._method_name

    @property
    def is_default_method(self):
        return self._is_default_method

    def pre_call_impl(self, caller, args):
        # This returns actual method, caller and args.
        args = self.args_impl(caller, args)
        actual_caller = caller
        if self.is_default_method:
            holder = actual_caller
        else:
            accessor = caller.accessor
            holder = accessor.get(caller, self.actual_caller)
            if holder is None:
                raise NetworkNothing("Method holder named'{}' not found.".format(self.actual_caller))
            if isinstance(holder, GenericValueHolder):
                holder = holder.value
            # print("***", type(holder), holder)
            if not isinstance(holder, NetworkInstance):
                raise NetworkError("Invalid method holder {}.".format(holder))
            actual_caller = caller
        # self.log.debug("**** pre  {}.get_method({},{})".format(holder, caller, self._method_name))
        callee = holder.get_callable(caller, self.method_name)
        # self.log.debug("done.")
        if callee is None:
            raise NetworkNotImplementationError("Method {}.{} not found.".format(holder, self.method_name))
        if isinstance(callee, NetworkMethod) and callee.instance_method and type(holder) is NetworkInstance:
            actual_args = [holder]
            actual_args.extend(args)
            actual_args = tuple(actual_args)
        else:
            actual_args = args
        return callee, actual_caller, actual_args

    def args_impl(self, caller, args, **kwargs):
        if not isinstance(caller, NetworkInstance):
            raise NetworkError("Invalid caller {}.".format(caller))
        args_converted = []
        for a in self.args:
            if isinstance(a, NetworkInstance):
                x = a
            elif isinstance(a, NetworkSymbol):
                x = caller.accessor.get(caller, a.symbol)
            elif isinstance(a, NetworkCallable):
                x = a(caller)
            elif isinstance(a, GenericValueHolder):
                x = a.value
            elif isinstance(a, GenericEvaluatee):
                x = a.evaluate(caller)
            else:
                x = a
            args_converted.append(x)
        return tuple(args_converted)

    def __repr__(self):
        args = ""
        if len(self.args) > 0:
            args = "{}".format(self.args[0])
            for a in self.args[1:]:
                args = "{},{}".format(args, a)
        return "{}({})".format(self.hierarchical_name, args)


class NetworkSubstituter(NetworkCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, var, callee, globally=False, security=NetworkInstance.STACK, overwrite=False):
        super().__init__(owner, args=tuple([callee]), cancel_stacking=True)
        self._var = var.symbol
        self._globally = globally
        self._security = security
        self._overwrite = overwrite

    def set_owner(self, owner):
        super().set_owner(owner)

    @property
    def var(self):
        return self._var

    @property
    def callee(self):
        return self.args[0]

    @property
    def globally(self):
        return self._globally

    @property
    def security(self):
        return self._security

    @property
    def overwrite(self):
        return self._overwrite

    def call_impl(self, caller, args, **kwargs):
        # print("*** substituting 1", self.var, "of", caller, "<=", self.callee)
        if self.globally:
            depth = 0
        else:
            depth = caller.deepest_stack_id(caller)
        self.log.debug("##### {0} <-- {1}".format(self.var, args[0]))
        #ret = caller.set_attribute(caller, self.var, args[0], depth=depth)
        caller.accessor.set(caller, self.var, args[0], security=self.security, globally=self.globally, overwrite=self.overwrite)
        # accessor.set(caller, self.var, args[0])
        # print("*** substituted", self.var, "of", caller, "<=", self.callee)

    def args_impl(self, caller, args, **kwargs):
        # print("*** caller", caller, self.args)
        # do nothing, since no argument preparation have to be done here.
        return super().args_impl(caller, self.args)

    def __repr__(self):
        return "{} <- {}".format(self.var, self.callee)


class NetworkBreak(NetworkCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    BREAK = "break"

    def __init__(self, owner):
        super().__init__(owner, closer=True)

    def call_impl(self, caller, args, **kwargs):
        caller.set_break_point(caller, self.BREAK)

    def args_impl(self, caller, args, **kwargs):
        return args

    def __repr__(self):
        return "break;"


class NetworkReturn(NetworkCallable, GenericValueHolder):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    RETURN = "return"

    def __init__(self, owner, callee):
        super().__init__(owner, closer=True)
        self._callee = callee

    @property
    def callee(self):
        return self._callee

    def call_impl(self, caller, args, **kwargs):
        accessor = caller.accessor
        if isinstance(self.callee, NetworkCallable):
            a = self.callee(caller)
        elif isinstance(self.callee, NetworkSymbol):
            a = accessor.get(caller, self.callee.symbol)
        elif isinstance(self.callee, GenericValueHolder):
            a = self.callee.value
        else:
            a = self.callee
        caller.set_break_point(caller, self.RETURN)
        self.log.debug("return {}".format(a))
        return a

    def __repr__(self):
        return "return {};".format(self.callee)


class NetworkSequentialStatements(NetworkCallable):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner):
        super().__init__(owner)
        self._statements = ()

    def set_owner(self, owner):
        super().set_owner(owner)
        for stmts in self._statements:
            if isinstance(stmts, GenericComponent):
                stmts.set_owner(owner)
            elif type(stmts) is list or type(stmts) is tuple:
                for s in stmts:
                    s.set_owner(owner)
            else:
                debug("Unexpected type '{}'", stmts)


    @property
    def statements(self):
        return self._statements

    def set_statements(self, stmt):
        self._statements = tuple(stmt)

    def append_statements(self, stmt):
        statements = list(self.statements)
        statements.append(stmt)
        self._statements = tuple(statements)

    def call_impl(self, caller, args, **kwargs):
        rtn = None
        for stmt in self.statements:
            rtn = stmt(caller, args)
            if caller.break_point is not None:
                caller.set_break_point(caller, None)
                return rtn
        return rtn

    def __repr__(self):
        stmts = ""
        for s in self.statements:
            stmts = "{} {}".format(stmts, s)
        return "{};".format(stmts)


class StatementParam(Enum):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    STEP = "step"
    CONDITIONAL_ARGS = "conditional-args"
    CONDITION_TRUELY_EVALUATED = "condition-truely-evaluated"
    STATEMENTS = "statements"
    OPTIONS = "options"
    REACHED = "reached"
    VISITED = "visited"
    IF = "if"
    ELIF = "elif"
    ELSE = "else"
    END = "end"


class ForeachStatement(NetworkSequentialStatements):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, var, fetchee):
        super().__init__(owner)
        self._var = var.symbol
        self._fetchee = fetchee.symbol

    @property
    def var(self):
        return self._var

    @property
    def fetchee(self):
        return self._fetchee

    def fetch(self, caller):
        accessor = caller.accessor
        fetchee = accessor.get(caller, self.fetchee)
        pos = accessor.get(caller, self.fetch_pos_name)
        pos = pos + 1
        accessor.set(caller, self.fetch_pos_name, pos)
        if pos < len(fetchee):
            accessor.set(caller, self.var, fetchee[pos])
            return True
        else:
            return False

    @property
    def fetch_pos_name(self):
        return "{}$pos".format(self.var)

    def prepare(self, caller):
        try:
            caller.accessor.set(caller, self.fetch_pos_name, -1)
            return True
        except NetworkReferenceError as ex:
            debug("preparation failed in ForeachStatement.")
            return False

    def call_impl(self, caller, args, **kwargs):
        self.log.debug("*** call_impl with self={}, caller={}, args={}".format(self, caller, args))
        rtn = None
        while self.fetch(caller):
            rtn = super().call_impl(caller, args)
            if caller.break_point is not None:
                caller.set_break_point(caller, None)
                break
        return rtn

    def args_impl(self, caller, args, **kwargs):
        # self.log.debug("*** args_impl with self={}, caller={}, args={}".format(self, caller, args))
        if not self.prepare(caller):
            rtn = NetworkReturnValue(caller, False, "fetching from {} failed.".format(self.fetchee))
            return rtn
        return args

    def __repr__(self):
        return "for({}:{})...".format(self.var, self.fetchee)


class ForallStatement(ForeachStatement):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, var, fetchee):
        super().__init__(owner, var, fetchee)

    def call_impl(self, caller, args, **kwargs):
        self.log.debug("*** call_impl with self={}, caller={}, args={}".format(self, caller, args))
        while self.fetch(caller):
            stmt = self.statements[0]
            rtn = stmt.evaluate(caller)
            # rtn = super().call_impl(caller, args)
            if not rtn:
                return False
        return True

    def __repr__(self):
        return "forall {} in {}: ...".format(self.var, self.fetchee)
        #return "forall {} in {}: {}".format(self.var, self.fetchee, self.statements)


class ExistsStatement(ForeachStatement):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, var, fetchee):
        super().__init__(owner, var, fetchee)

    def call_impl(self, caller, args, **kwargs):
        self.log.debug("*** call_impl with self={}, caller={}, args={}".format(self, caller, args))
        while self.fetch(caller):
            stmt = self.statements[0]
            rtn = stmt.evaluate(caller)
            if rtn:
                return True
        return False

    def __repr__(self):
        return "exists {} in {}: ...".format(self.var, self.fetchee)
        #return "exists {} in {}: {}".format(self.var, self.fetchee, self.statements)


class WhileStatement(NetworkSequentialStatements):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, cond):
        super().__init__(owner)
        self._cond = cond

    @property
    def condition(self):
        return self._cond

    def call_impl(self, caller, args, **kwargs):
        rtn = None

        while self.condition.evaluate(caller):
            rtn = super().call_impl(caller, args)
            if caller.break_point is not None:
                caller.set_break_point(caller, None)
                break
        return rtn

    def __repr__(self):
        return "while(..){..}"
        #return "while({}){}".format(self._cond, self.statements)


class IfElifElseStatement(NetworkSequentialStatements):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner):
        super().__init__(owner)
        self._closed = False
        self._conditions = []

    @property
    def closed(self):
        return self._closed

    @property
    def conditions(self):
        return self._conditions

    def append_if_elif_statements(self, condition, statements):
        if self.closed:
            raise NetworkError("{} illegal attempt. statement already closed.".format(self))
        else:
            self.conditions.append(condition)
            self.append_statements(statements)

    def append_else_statements(self, statements):
        if self.closed:
            raise NetworkError("{} illegal attempt. statement already closed.".format(self))
        else:
            condition = TrueEvaluatee(None)
            self.conditions.append(condition)
            self.append_statements(statements)
            self._closed = True

    def call_impl(self, caller, args, **kwargs):
        rtn = None
        for cond, stmt in zip(self.conditions, self.statements):
            if cond.evaluate(caller):
                for s in stmt:
                    rtn = s(caller)
                    if caller.break_point is not None:
                        caller.set_break_point(caller, None)
                        return rtn
                return rtn
        return rtn

    def __repr__(self):
        return "if(..){..}elif(..){..}else{..}"
        # repr = "if ({}) {}".format(self._conditions[0], len(self.statements[0]))
        # for c, s in zip(self.conditions[1:], self.statements[1:]):
        #     if c == self.conditions[len(self.conditions)-1]:
        #         _repr = "else {}".format(s)
        #     else:
        #         _repr = "elif ({}) {}".format(c, s)
        #     repr = "{} {}".format(repr, _repr)
        # return repr


class MethodAvator(NetworkMethod):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, method_owner, signature, callee_owner, callee_entity, extra_args, globally=True):
        super().__init__(method_owner, signature, args=None, globally=globally)
        self._callee_owner = callee_owner
        self._callee_entity = callee_entity
        self._extra_args = extra_args

    def call_impl(self, caller, args, **kwargs):
        if self._callee_owner is None:
            return self._callee_entity(caller, args, self._extra_args)
        else:
            return self._callee_entity(self._callee_owner, caller, args, self._extra_args)

    def args_impl(self, caller, args, **kwargs):
        new_args = []
        for a in args:
            if isinstance(a, GenericValueHolder):
                new_args.append(a.value)
            else:
                new_args.append(a)
        return new_args

    def __repr__(self):
        return "{}.{}:wrapper(actual owner:{})".format(self.owner.signature, self.signature, self._callee_owner)


class BoundMethodAvator(NetworkMethod):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, method_owner, signature, callee_owner, method_entity, globally=True):
        super().__init__(method_owner, signature, args=None, globally=globally)
        self._callee_owner = callee_owner
        self._method_entity = method_entity

    def call_impl(self, caller, args, **kwargs):
        args_str = ""
        if len(args) != 0:
            args_str = "a[0]"
            for i in range(1, len(args)):
                args_str = "{},a[{}]".format(args_str, i)
        eqn = "lambda o, a: o.{}({})".format(self.signature, args_str)
        func = eval(eqn)
        return func(caller, args)

    def args_impl(self, caller, args, **kwargs):
        new_args = []
        for a in args:
            if isinstance(a, GenericValueHolder):
                new_args.append(a.value)
            else:
                new_args.append(a)
        return new_args

    def __repr__(self):
        return "{}.{}:wrapper(actual owner:{})".format(self.owner.signature, self.signature, self._callee_owner)


class WrappedAccessor(NetworkMethod):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, accessor_owner, signature, actual_owner, accessor, globally=True, extra_args=None, **kwargs):
        super().__init__(accessor_owner, signature, args=None, globally=globally)
        self._actual_owner = actual_owner
        self._accessor = accessor
        self._extra_args = extra_args
        self._kwargs = {}
        for k in kwargs.keys():
            self._kwargs[k] = kwargs[k]

    @property
    def accessor_owner(self):
        return self.owner

    @property
    def accessor(self):
        return self._accessor

    @property
    def actual_owner(self):
        return self._actual_owner

    @property
    def extra_args(self):
        return self._extra_args

    @property
    def kwargs(self):
        return self._kwargs

    def call_impl(self, caller, args, **kwargs):
        rtn = self.accessor(self.accessor_owner, caller, self.actual_owner, args, self.extra_args)
        return rtn

    def args_impl(self, caller, args, **kwargs):
        new_args = []
        for a in args:
            if isinstance(a, GenericValueHolder):
                new_args.append(a.value)
            else:
                new_args.append(a)
        return new_args

    def __repr__(self):
        return "WrappedAccessor to {}.{}".format(self.actual_owner, self.signature)


class StrictWrappedAccessor(WrappedAccessor):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, accessor_owner, signature, actual_owner, accessor, args_validator=None, globally=True, extra_args=None, **kwargs):
        super().__init__(accessor_owner, signature, actual_owner=actual_owner, accessor=accessor,
                         globally=globally, extra_args=extra_args, kwargs=kwargs)
        self._args_validator = args_validator

    @property
    def args_validator(self):
        return self._args_validator

    # No need to implement, due to identical to super-class's method.
    # def call_impl(self, caller, args, **kwargs):
    #     caller_args = args
    #     result = self.accessor(self.accessor_owner, caller, self.actual_owner, caller_args, self.extra_args)
    #     return result

    def args_impl(self, caller, args, **kwargs):
        if self.args_validator is None:
            return args
        caller_args = args
        result = self.args_validator(self.accessor_owner, caller, self.actual_owner, caller_args, self.extra_args)
        if not isinstance(result, NetworkReturnValue):
            return result
        return result.value
        # elif result:
        #     rtn = NetworkReturnValue(caller_args, True, "")
        # else:
        #     rtn = NetworkReturnValue(caller_args, False, "argument check error")
        # return rtn

    def __repr__(self):
        return "StrictWrappedAccessor to {}.{}".format(self.actual_owner, self.signature)


class ExtensibleWrappedAccessor(StrictWrappedAccessor):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, accessor_owner, signature, actual_owner, accessor, args_validator=None, globally=True,
                 extra_args=(), help_text="", preprocesses=(), postprocesses=(), **kwargs):
        super().__init__(accessor_owner, signature, actual_owner=actual_owner, accessor=accessor,
                         args_validator=args_validator,
                         globally=globally, extra_args=extra_args, kwargs=kwargs)
        self._preprocesses = preprocesses
        self._postprocesses = postprocesses
        self._help_text = help_text

    @property
    def preprocesses(self):
        return self._preprocesses

    @property
    def postprocesses(self):
        return self._postprocesses

    @property
    def help_text(self):
        return self._help_text

    def args_impl(self, caller, args, **kwargs):
        if self.args_validator is None:
            return super().args_impl(caller, args)
        caller_args = args
        result = self.args_validator(self.accessor_owner, caller, self.actual_owner, caller_args, self.extra_args)
        if not isinstance(result, NetworkReturnValue):
            return result
        return result.value

    def call_impl(self, caller, args, **kwargs):
        if len(args) > 0 and isinstance(args[0], CommandOption) and args[0].name == "help" and self.help_text is not None:
            caller.print(self.help_text)
            return None
        args = self.args_impl(caller, args)

        # Checks pre-condition.
        for pre in self.preprocesses:
            p = pre(self.accessor_owner, caller, self.actual_owner, args, self.extra_args)
            if not p:
                return p
        # Calls accessor.
        rtn = self.accessor(self.accessor_owner, caller, self.actual_owner, args, self.extra_args)
        # Checks post-condition.
        for post in self.postprocesses:
            p = post(self.accessor_owner, caller, self.actual_owner, args, self.extra_args)
            if not p:
                break
        return rtn

    def __repr__(self):
        return "{}".format(self.signature)


class MethodWrapper(NetworkMethod):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, method_owner, signature, default_caller, callee_entity, args):
        super().__init__(method_owner, signature, args, [])
        self._default_caller = default_caller
        self._callee_entity = callee_entity

    def call_impl(self, caller, args, **kwargs):
        if self._default_caller is not None:
            caller = self._default_caller
        return self._callee_entity(caller, args)

    def __repr__(self):
        return "{}:wrapper".format(self.signature)


class NetworkGenericWorld(NetworkClass):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, signature, super_class, embedded=()):
        super().__init__(owner, signature, super_class, embedded=embedded)
        self._resolvers = []
        self._bound_methods = []

    def append_resolver(self, resolver):
        self._resolvers.append(resolver)

    def remove_resolver(self, resolver):
        self._resolvers.remove(resolver)

    def deserialize(self, serializer: NetworkSerializer):
        try:
            serializer.open(for_read=True)
            # script = serializer.deserialize()
            serializer.close()
        except Exception as ex:
            if serializer.opened:
                serializer.close()
            raise ex

    def serialize(self, serializer: NetworkSerializer):
        try:
            serializer.open(for_read=False, for_write=True)
            for clazz in self.classes:
                doc = clazz.document
                doc.serizlize(serializer)
        except Exception as ex:
            if serializer.opened:
                serializer.close()
            raise ex

    def get_method(self, caller, sig):
        return super().get_method(caller, sig)
        # FIXME initially, mediator was planned to decide adequate methods.
        #       but now instance itself  decides adequateness.
        # m = super().get_method(caller, sig)
        # if m is not None:
        #     return m
        # for r in self._resolvers:
        #     m = r.get_method(caller, sig)
        #     if m is not None:
        #         return m
        # return None

    def get_bound_method(self, obj, args):
        # FIXME this method should not be called.
        self.log.debug("*** GET_BOUND_METHOD IS DEPRECATED. ***")
        for m in inspect.getmembers(obj, inspect.ismethod):
            if m[0] == args:
                return m[1]
        return None

    def get_bound_methods_info(self, obj, args):
        # FIXME this method should not be called.
        self.log.debug("*** GET_BOUND_METHOD IS DEPRECATED. ***")
        I = []
        for m, _ in inspect.getmembers(obj, inspect.ismethod):
            # if m in args:
            I.append(m)
        return I

    # def implement_bound_methods(self, obj, args):
    #     R = []
    #     for n, m in inspect.getmembers(obj, inspect.ismethod):
    #         if n in args:
    #             # callee_entity = lambda c, caller, caller_args, generator_args: generator_args[0](caller, caller_args)
    #             # method = BoundMethodAvator(obj, n, obj, callee_entity, [m], globally=True)
    #             method = BoundMethodAvator(obj, n, obj, m, globally=True)
    #             obj.declare_method(method, globally=True)
    #             R.append(n)
    #     return R


class NetworkMethodWrapper(NetworkMethod):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, owner, swapped_context, signature, func, args):
        super().__init__(owner, signature, args, [])
        self._swapped_context = swapped_context
        self._func = func

    def set_context(self, context):
        self._swapped_context = context

    @property
    def swapped_context(self):
        return self._swapped_context

    @property
    def document(self):
        return self._document

    def call_impl(self, caller, args, **kwargs):
        return self._func(self._swapped_context, args)


class ResolverMediator:

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, mediator: NetworkInvoker):
        self._mediator = mediator

    @property
    def mediator(self) -> NetworkInvoker:
        return self._mediator


class ToplevelInstance(NetworkInstance):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, signature):
        super().__init__(None, 0, None, None)
        self._signature = signature
        self._mediators = []

    @property
    def signature(self):
        return self._signature

    @property
    def mediators(self):
        return self._mediators

    def append_mediator(self, med):
        self._mediators.append(med)

    def remove_mediator(self, med):
        self._mediators.remove(med)


class NetworkSecurityPolicy(GenericDescription):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self, name, is_symbol=False):
        super().__init__(name, is_symbol)
        self._desc = None
        self._map = {}

    @property
    def description(self):
        return self._desc

    def set_descriptoin(self, desc):
        self._desc = desc

    def encode_description(self):
        if self.description is None:
            return NetworkReturnValue(False, False, "no description")
        pos = 0
        skip_pat = r"\n*(\s*(//.*|)\n)*"
        # desc_pat = r"\s*(?P<start>[a-zA-Z]+0-9 ]*[a-zA-Z0-9]+)::\s*\n(?P<desc>.*)\n\s*::(?P<end>[a-zA-Z]+0-9 ]*[a-zA-Z0-9]+)(\s*\n)*"
        desc_pat = r"\s*(?P<type>(property|method))\s+(?P<member>)\s*::\s*(?P<desc>(private|protected|public)(\s*\n)*"
        while True:
            m, d, next_pos = GU.pattern_iter(skip_pat, self.description[pos:])

            if m is not None:
                pos = next_pos
                continue
            m, d, next_pos = GU.pattern_iter(desc_pat, self.description[pos:])
            if m is None:
                break
            t = d['type']
            m = d['member']
            desc = d['desc']
            self._map[m] = {}
            self._map[m]["type"] = t
            self._map[m]["desc"] = desc


class MethodArmer(NetworkInstance):

    log = log4p.GetLogger(logger_name=__name__, config=config.get_log_config()).logger

    def __init__(self):
        super().__init__(None, 0, None)
        self._conf = GU.read_json("arm.conf")

    def _arm_method(self, sig, instance, generator, manager, method_config):
        globally = method_config["globally"]
        m = None
        if "equation" in method_config.keys():
            eqn = method_config["equation"]
            eqn = eval(eqn)
            if "help-text" in method_config.keys():
                help_text = method_config["help-text"]
            else:
                help_text = "Help text is not prepared."
            # arguments validator
            if "args_validator" in method_config.keys():
                args_validator = eval(method_config["args_validator"])
            else:
                args_validator = None
            # pre processes
            preprocess = []
            if "preprocess" in method_config.keys():
                config = method_config["preprocess"]
                for p in sorted(config.keys()):
                    preprocess.append(eval(config[p]))
            preprocess = tuple(preprocess)
            # post processes
            postprocess = []
            if "postprocess" in method_config.keys():
                config = method_config["postprocess"]
                for p in sorted(config.keys()):
                    postprocess.append(eval(config[p]))
            postprocess = tuple(postprocess)
            # Creates wrapped accessor
            m = ExtensibleWrappedAccessor(instance, sig, manager, eqn,
                                          help_text=help_text,
                                          args_validator=args_validator,
                                          preprocesses=preprocess,
                                          postprocesses=postprocess,
                                          globally=globally)
        elif "script" in method_config.keys():
            interpreter = instance.get_method(instance, "interpret")
            if interpreter is None:
                self.log.debug("*** INSTANCE '{}' HAS NO INTERPRETER. SCRIPT CANNOT READ.".format(instance))
            else:
                script = method_config["script"]
                m = interpreter(self, [script])
        elif "script-file" in method_config.keys():
            with open(method_config["script-file"]) as f:
                script = f.read()
            interpreter = instance.get_method(instance, "interpret")
            if interpreter is None:
                self.log.debug("*** INSTANCE '{}' HAS NO INTERPRETER. SCRIPT CANNOT READ.".format(instance))
            else:
                m = interpreter(self, [script])
        else:
            m = None
            self.log.debug("Invalid configuration description for method {}".format(sig))
        if m is not None:
            instance.register_method(instance, sig, m, depth=0, overwrite=None)

    def _arm_methods(self, instance, generator, manager, group, dispatcher):
        for g in group:
            for d in dispatcher[g]:
                method_config = dispatcher[g][d]
                self._arm_method(d, instance, generator, manager, method_config)

    def _arm_default_interpreter(self, instance, generator, manager):
        interpreter = IM.get_interpreter(instance, generator, manager)
        if interpreter is None:
            self.log.debug("*** INSTANCE '{}' HAS NO INTERPRETER. FOR THIS, SCRIPT CANNOT BE READ PRESENTLY.".format(instance))
        else:
            instance.register_method(instance, interpreter.signature, interpreter)

    def arm_default_methods(self, instance, generator, manager):
        self._arm_default_interpreter(instance, generator, manager)
        group = self._conf["method-group"]
        dispatcher = self._conf["method-dispatcher"]
        for g in group.keys():
            if g in instance.embedded_attrib(instance):
                self._arm_methods(instance, generator, manager, group[g], dispatcher)

    def arm_methods(self, instance, generator, manager):
        # FIXME current implementation is same to arm_default_methods.
        self.arm_default_methods(instance, generator, manager)
        # self._arm_default_interpreter(instance, generator, manager)
        # group = self._conf["method-group"]
        # dispatcher = self._conf["method-dispatcher"]
        # for g in group.keys():
        #     if g in instance.embedded_attrib(instance):
        #         self._arm_methods(instance, generator, manager, group[g], dispatcher)

    def arm_method(self, sig, instance, generator, manager, method_config):
        self._arm_method(sig, instance, generator, manager, method_config)


# set default armer
set_armer(MethodArmer())

